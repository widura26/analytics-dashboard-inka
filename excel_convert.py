"""Pembaca file Excel -> grid sel in-memory (semua sheet, tanpa konversi CSV).

Dipakai alur ETL: begitu user mengunggah file .xlsx/.xls, SELURUH sheet dibaca
lewat pandas menjadi grid (daftar baris × kolom sel string) — format yang sama
dengan hasil pembacaan CSV — sehingga tahap rekap & cleaning bekerja sama persis
baik untuk file CSV maupun Excel.

Grid dipertahankan apa adanya (header=None, tanpa inferensi nama kolom) supaya
deteksi header multi-baris pada rekap BOM dan pembacaan SAP bekerja sama persis
seperti pada file CSV. Sel angka bulat ditulis tanpa ".0" (kode material
1000123 tidak boleh berubah jadi "1000123.0"), dan sel tanggal diformat ulang
memakai format tanggal standar aplikasi (dd.mm.yyyy).

PEMBACAAN FORMAT BARIS (khusus .xlsx/.xlsm): selain nilai, tiap sheet juga
dipindai untuk menemukan baris yang DITANDAI orang di Excel sebagai "sudah
dihapus/tidak dipakai":

  - tersembunyi/collapse  -> ws.row_dimensions[r].hidden
  - dicoret (strikethrough)-> cell.font.strike
  - blok gelap/hitam      -> cell.fill solid dengan semua kanal RGB <= 63

Hasilnya dict {'baris 0-based': [alasan, ...]} pada kunci 'tandai' tiap sheet,
yang diteruskan ke bom_rekap.rekap_grid(abaikan=...) supaya baris itu TIDAK
ikut rekap (lihat spec 2026-10-20-lewati-baris-berformat-excel.md). Pandas
tidak punya akses ke format ini, jadi dibuka sekali lagi dengan openpyxl mode
normal (mode read-only tidak membawa row_dimensions). File .xls/.xlsb tidak
mendukung pembacaan format -> 'tandai' kosong (perilaku lama).
"""
import io
from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles.colors import COLOR_INDEX

# Format tanggal standar aplikasi (sama dengan SAP_DATE_FORMAT).
DEFAULT_DATE_FORMAT = "%d.%m.%Y"

# Ekstensi yang dikenali beserta engine pembacanya.
ENGINE_BY_SUFFIX = {
    ".xlsx": "openpyxl",
    ".xlsm": "openpyxl",
    ".xls": "xlrd",
    ".xlsb": "pyxlsb",
}
EXCEL_SUFFIXES = tuple(ENGINE_BY_SUFFIX)

# Alasan baris ditandai (dipakai sebagai nilai dict 'tandai' + laporan rekap).
ALASAN_HIDDEN = "tersembunyi/collapse"
ALASAN_STRIKE = "dicoret"
ALASAN_FILL = "blok gelap"

# Blok dianggap "hitam/gelap" bila SEMUA kanal RGB <= ini (63 = #3F).
# Sengaja longgar untuk hitam pekat, tapi menolak warna gelap biasa
# (mis. navy #1F3864 punya kanal B=100 -> bukan penanda).
BATAS_GELAP = 63

# Batas kolom dipindai (grid BOM jauh lebih sempit; mencegah scan ratusan
# ribu kolom kalau ada format sisa di kolom XFD).
MAKS_KOLOM_PINDAI = 64


def file_suffix(filename: str) -> str:
    """Ekstensi huruf kecil tanpa titik: 'bom.csv' -> 'csv'."""
    return Path(filename).suffix.lower().lstrip(".")


def file_stem(filename: str) -> str:
    """Nama file tanpa ekstensi: 'bom_proyek.csv' -> 'bom_proyek'."""
    return Path(filename).stem


def is_excel_file(filename: str) -> bool:
    """True bila nama file berekstensi Excel yang dikenali."""
    return Path(filename).suffix.lower() in ENGINE_BY_SUFFIX


def read_excel_sheets(
    raw: bytes,
    filename: str,
    date_format: str = DEFAULT_DATE_FORMAT,
):
    """Baca SEMUA sheet file Excel menjadi grid sel (tanpa lewat CSV).

    Mengembalikan daftar dict per sheet:
      {"label": nama sheet, "grid": list[list[str]], "rows": int, "cols": int,
       "tandai": {baris 0-based: [alasan, ...]}}

    'tandai' berisi baris yang ditandai di Excel (tersembunyi/collapse,
    dicoret, blok gelap) — hanya untuk format yang bisa dibaca
    (.xlsx/.xlsm); file lain -> {} (lihat modul docstring).
    """
    suffix = Path(filename).suffix.lower()
    engine = ENGINE_BY_SUFFIX.get(suffix)
    if engine is None:
        raise ValueError(
            f"Format '{suffix or filename}' bukan Excel yang didukung "
            f"({', '.join(sorted(ENGINE_BY_SUFFIX))})."
        )

    try:
        excel = pd.ExcelFile(io.BytesIO(raw), engine=engine)
    except ImportError as exc:
        raise ValueError(
            f"Pembaca format '{suffix}' ('{engine}') belum terpasang di "
            f"server aplikasi. {exc}"
        ) from exc
    except Exception as exc:
        raise ValueError(f"Gagal membuka file Excel: {exc}") from exc

    # Pemindaian format baris terpisah dari pandas (pandas tidak membaca
    # format sel). Gagal di sini tidak boleh membatalkan pembacaan nilai.
    tandai = _baca_tandai_baris(raw) if engine == "openpyxl" else {}

    hasil = []
    try:
        for title in excel.sheet_names:
            sheet_df = excel.parse(title, header=None)
            grid = [
                [_cell_to_text(value, date_format) for value in row]
                for row in sheet_df.itertuples(index=False, name=None)
            ]
            hasil.append({
                "label": title,
                "grid": grid,
                "rows": int(len(sheet_df)),
                "cols": int(len(sheet_df.columns)),
                "tandai": tandai.get(title, {}),
            })
    except Exception as exc:
        raise ValueError(f"Gagal membaca sheet Excel: {exc}") from exc
    finally:
        excel.close()

    if not hasil:
        raise ValueError("File Excel tidak berisi sheet apa pun.")

    return hasil


# ============================================================
# PEMINDAIAN FORMAT: BARIS YANG DITANDAI ORANG DI EXCEL
# ============================================================

def _baca_tandai_baris(raw: bytes) -> dict:
    """Pindai format tiap sheet: baris tersembunyi/collapse, sel dicoret,
    dan sel berblok gelap/hitam.

    Mengembalikan {nama_sheet: {baris 0-based: [alasan, ...]}} — indeks
    0-based persis seperti posisi baris pada grid hasil _cell_to_text.
    Gagal membuka workbook (file rusak, format tak terduga) -> {} supaya
    pembacaan nilai tetap berjalan seperti biasa.
    """
    try:
        workbook = load_workbook(io.BytesIO(raw), data_only=True)
    except Exception:
        return {}

    hasil = {}
    try:
        for worksheet in workbook.worksheets:
            try:
                hasil[worksheet.title] = _pindai_tandai_sheet(worksheet)
            except Exception:
                hasil[worksheet.title] = {}
    finally:
        workbook.close()

    return hasil


def _pindai_tandai_sheet(worksheet) -> dict:
    """Baris ditandai dalam satu worksheet: {baris 0-based: [alasan, ...]}.

    Tiga penanda yang dikenali (keputusan user 2026-10-20):
      - baris disembunyikan / tergabung dalam group collapse
      - sel memakai font dicoret (strikethrough)
      - sel diblok warna gelap/hitam (fill solid nyaris hitam)

    Baris header tidak dikecualikan di sini — pemotongan baris terjadi di
    bom_rekap.rekap_grid, yang hanya memakai tanda untuk baris DATA.
    """
    tandai = {}

    def tandai_baris(idx: int, alasan: str):
        daftar = tandai.setdefault(idx, [])
        if alasan not in daftar:
            daftar.append(alasan)

    # 1) baris tersembunyi / collapse (row_dimensions hanya terisi di
    #    mode normal; memang tidak tersedia di mode read-only)
    for nomor, dimensi in worksheet.row_dimensions.items():
        if int(nomor) >= 1 and getattr(dimensi, "hidden", False):
            tandai_baris(int(nomor) - 1, ALASAN_HIDDEN)

    # 2) sel dicoret / diblok gelap. Batas kolom dibatasi supaya bekas
    #    format di kolom jauh (XFD) tidak memperlambat scan.
    maks_col = min(worksheet.max_column, MAKS_KOLOM_PINDAI)
    for baris in worksheet.iter_rows(
        min_row=1, max_row=worksheet.max_row, max_col=maks_col
    ):
        idx = baris[0].row - 1
        for cell in baris:
            if getattr(cell.font, "strike", False):
                tandai_baris(idx, ALASAN_STRIKE)
            if (
                cell.fill.patternType == "solid"
                and _warna_gelap(cell.fill.fgColor)
            ):
                tandai_baris(idx, ALASAN_FILL)

    return tandai


def _warna_gelap(color) -> bool:
    """True bila warna fill nyaris hitam (semua kanal RGB <= BATAS_GELAP).

    Mendukung warna tipe rgb, indexed (palet Excel), dan theme (hanya
    theme 1 = dk1/Hitam bawaan Excel; accent tidak dihitung supaya tabel
    berwarna gelap-biasa tidak ikut terlewat).
    """
    if color is None:
        return False

    tipe = getattr(color, "type", None)

    if tipe == "theme":
        return getattr(color, "theme", None) == 1 and not getattr(
            color, "tint", 0
        )

    if tipe == "indexed":
        nomor = getattr(color, "indexed", None)
        try:
            if nomor is None or not 0 <= int(nomor) < len(COLOR_INDEX):
                return False
        except (TypeError, ValueError):
            return False
        rgb = str(COLOR_INDEX[int(nomor)])
    else:
        rgb = str(getattr(color, "rgb", "") or "")

    rgb = rgb[-6:]  # buang prefix alpha "AA" pada AARRGGBB
    if len(rgb) != 6:
        return False

    try:
        kanal = [int(rgb[i:i + 2], 16) for i in (0, 2, 4)]
    except ValueError:
        return False

    return max(kanal) <= BATAS_GELAP


def _cell_to_text(value, date_format: str = DEFAULT_DATE_FORMAT) -> str:
    """Ubah satu sel Excel menjadi teks sel seperti pada file CSV."""
    if pd.isna(value):
        return ""

    if isinstance(value, bool):
        return str(value)

    # Termasuk pd.Timestamp (turunan datetime) & date.
    if isinstance(value, (datetime, date)):
        return value.strftime(date_format)

    if isinstance(value, (int, float)):
        number = float(value)
        if number.is_integer():
            return str(int(number))
        return str(number)

    return str(value)
